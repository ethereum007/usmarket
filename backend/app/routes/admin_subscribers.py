"""Read-only Alpha Radar subscriber dashboard endpoints."""

from __future__ import annotations

import logging
from io import BytesIO
from datetime import datetime, timedelta, timezone
from typing import Any, Literal

from fastapi import APIRouter, Depends, HTTPException, Query, Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from supabase import Client

from app.dependencies import get_supabase, require_admin_request

logger = logging.getLogger("alpha_radar.routes.admin_subscribers")

router = APIRouter(prefix="/api/admin", tags=["admin"])

MAX_ADMIN_USERS = 1_000


def _value(source: Any, key: str, default: Any = None) -> Any:
    """Read a field from either a Supabase model or a dict used in tests."""
    if isinstance(source, dict):
        return source.get(key, default)
    return getattr(source, key, default)


def _iso(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    return str(value)


def _as_datetime(value: Any) -> datetime | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        parsed = value
    else:
        try:
            parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        except (TypeError, ValueError):
            return None
    return parsed.replace(tzinfo=parsed.tzinfo or timezone.utc)


def _masked_reference(value: Any) -> str | None:
    reference = str(value or "").strip()
    if not reference:
        return None
    if len(reference) <= 10:
        return reference
    return f"{reference[:5]}…{reference[-5:]}"


def _member_name(metadata: Any, email: str | None) -> str | None:
    safe_metadata = metadata if isinstance(metadata, dict) else {}
    for key in ("full_name", "name", "display_name"):
        value = safe_metadata.get(key)
        if isinstance(value, str) and value.strip():
            return value.strip()

    first_name = safe_metadata.get("first_name")
    last_name = safe_metadata.get("last_name")
    name = " ".join(
        part.strip()
        for part in (str(first_name or ""), str(last_name or ""))
        if part and str(part).strip()
    )
    if name:
        return name
    return (email or "").split("@", maxsplit=1)[0] or None


def _mobile_number(user: Any, metadata: Any) -> str | None:
    phone = _value(user, "phone")
    if isinstance(phone, str) and phone.strip():
        return phone.strip()

    safe_metadata = metadata if isinstance(metadata, dict) else {}
    for key in ("mobile", "phone", "phone_number"):
        value = safe_metadata.get(key)
        if isinstance(value, str) and value.strip():
            return value.strip()
    return None


def _effective_status(subscription: dict[str, Any], now: datetime) -> tuple[str, bool]:
    plan = str(subscription.get("plan") or "free").lower()
    stored_status = str(subscription.get("status") or "active").lower()
    end_at = _as_datetime(subscription.get("current_period_end"))
    expired = plan == "premium" and end_at is not None and end_at <= now

    if expired or stored_status == "expired":
        return "expired", True
    if plan == "premium" and stored_status == "active":
        return "premium", False
    return "free", False


def _payment_status(subscription: dict[str, Any], membership_status: str) -> str:
    payment_id = subscription.get("razorpay_customer_id")
    order_id = subscription.get("razorpay_subscription_id")
    if payment_id:
        return "paid"
    if order_id:
        return "pending"
    if membership_status in {"premium", "expired"}:
        return "complimentary"
    return "not_started"


def _build_member(user: Any, subscription: dict[str, Any], now: datetime) -> dict[str, Any]:
    email = _value(user, "email")
    metadata = _value(user, "user_metadata", {})
    membership_status, is_expired = _effective_status(subscription, now)
    payment_status = _payment_status(subscription, membership_status)
    payment_id = subscription.get("razorpay_customer_id")
    order_id = subscription.get("razorpay_subscription_id")
    payment_date = _iso(subscription.get("current_period_start"))

    return {
        "id": str(_value(user, "id")),
        "name": _member_name(metadata, email),
        "email": email,
        "mobile": _mobile_number(user, metadata),
        "joinedAt": _iso(_value(user, "created_at")),
        "lastSignInAt": _iso(_value(user, "last_sign_in_at")),
        "membership": {
            "plan": str(subscription.get("plan") or "free").lower(),
            "status": membership_status,
            "isExpired": is_expired,
            "validFrom": payment_date,
            "validUntil": _iso(subscription.get("current_period_end")),
        },
        "payment": {
            "status": payment_status,
            "paidAt": payment_date if payment_status == "paid" else None,
            "amountPaise": 99_900 if payment_status == "paid" else None,
            "paymentReference": _masked_reference(payment_id),
            "orderReference": _masked_reference(order_id),
        },
    }


def _sort_datetime(value: str | None) -> datetime:
    return _as_datetime(value) or datetime.min.replace(tzinfo=timezone.utc)


def _matches_status(member: dict[str, Any], requested: str) -> bool:
    if requested == "all":
        return True
    status = member["membership"]["status"]
    if requested == "premium":
        return status == "premium"
    if requested == "expired":
        return status == "expired"
    return status == "free"


def _matches_search(member: dict[str, Any], query: str) -> bool:
    if not query:
        return True
    haystack = " ".join(
        str(member.get(field) or "") for field in ("name", "email", "mobile")
    ).casefold()
    return query.casefold() in haystack


def _load_members(sb: Client, now: datetime) -> list[dict[str, Any]]:
    subscription_result = (
        sb.table("subscriptions")
        .select(
            "user_id,plan,status,razorpay_subscription_id,razorpay_customer_id,"
            "current_period_start,current_period_end"
        )
        .limit(MAX_ADMIN_USERS)
        .execute()
    )
    subscriptions = subscription_result.data or []
    auth_users = sb.auth.admin.list_users(page=1, per_page=MAX_ADMIN_USERS)
    subscriptions_by_user = {
        str(row.get("user_id")): row
        for row in subscriptions
        if row.get("user_id")
    }
    return [
        _build_member(user, subscriptions_by_user.get(str(_value(user, "id")), {}), now)
        for user in auth_users
        if _value(user, "id")
    ]


def _filter_and_sort_members(
    members: list[dict[str, Any]],
    *,
    query: str,
    status_filter: str,
    sort: str,
) -> list[dict[str, Any]]:
    filtered = [
        member
        for member in members
        if _matches_status(member, status_filter) and _matches_search(member, query.strip())
    ]
    sort_key = {
        "recent_activity": lambda member: _sort_datetime(member["lastSignInAt"]),
        "joined": lambda member: _sort_datetime(member["joinedAt"]),
        "valid_until": lambda member: _sort_datetime(member["membership"]["validUntil"]),
        "payment": lambda member: _sort_datetime(member["payment"]["paidAt"]),
    }[sort]
    filtered.sort(key=sort_key, reverse=True)
    return filtered


def _build_members_workbook(members: list[dict[str, Any]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Alpha Radar Members"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:B{max(1, len(members) + 1)}"
    sheet.column_dimensions["A"].width = 32
    sheet.column_dimensions["B"].width = 44

    header_fill = PatternFill("solid", fgColor="009B72")
    for column, value in enumerate(("Name", "Email"), start=1):
        cell = sheet.cell(row=1, column=column, value=value)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="left")

    for row, member in enumerate(members, start=2):
        sheet.cell(row=row, column=1, value=str(member.get("name") or ""))
        sheet.cell(row=row, column=2, value=str(member.get("email") or ""))

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


@router.get("/subscribers")
async def list_subscribers(
    q: str = Query("", max_length=100),
    status_filter: Literal["all", "premium", "free", "expired"] = Query("all", alias="status"),
    sort: Literal["recent_activity", "joined", "valid_until", "payment"] = Query("recent_activity"),
    page: int = Query(1, ge=1),
    limit: int = Query(25, ge=10, le=100),
    _admin: dict = Depends(require_admin_request),
    sb: Client = Depends(get_supabase),
) -> dict[str, Any]:
    """Return the operational member view for the Alpha Login admin panel."""
    now = datetime.now(timezone.utc)
    try:
        members = _load_members(sb, now)
    except Exception as exc:
        logger.exception("[admin] failed to load Alpha Radar subscribers")
        raise HTTPException(
            status_code=500,
            detail="Unable to load subscriber dashboard. Please retry shortly.",
        ) from exc

    filtered = _filter_and_sort_members(
        members,
        query=q,
        status_filter=status_filter,
        sort=sort,
    )

    active_premium = sum(member["membership"]["status"] == "premium" for member in members)
    expired_premium = sum(member["membership"]["status"] == "expired" for member in members)
    free_members = sum(member["membership"]["status"] == "free" for member in members)
    logins_30d_cutoff = now - timedelta(days=30)
    recent_logins = sum(
        (_as_datetime(member["lastSignInAt"]) or datetime.min.replace(tzinfo=timezone.utc)) >= logins_30d_cutoff
        for member in members
        if member["lastSignInAt"]
    )

    total = len(filtered)
    offset = (page - 1) * limit
    return {
        "success": True,
        "generatedAt": now.isoformat(),
        "summary": {
            "totalUsers": len(members),
            "activePremium": active_premium,
            "freeUsers": free_members,
            "expiredPremium": expired_premium,
            "loginsLast30Days": recent_logins,
        },
        "pagination": {
            "page": page,
            "limit": limit,
            "total": total,
            "totalPages": max(1, (total + limit - 1) // limit),
        },
        "members": filtered[offset : offset + limit],
        "dataNote": (
            "Payment values reflect the latest known subscription payment. "
            "Alpha Radar does not yet keep a separate historical payment ledger."
        ),
    }


@router.get("/subscribers/export")
async def export_subscribers(
    q: str = Query("", max_length=100),
    status_filter: Literal["all", "premium", "free", "expired"] = Query("all", alias="status"),
    sort: Literal["recent_activity", "joined", "valid_until", "payment"] = Query("recent_activity"),
    _admin: dict = Depends(require_admin_request),
    sb: Client = Depends(get_supabase),
) -> Response:
    """Download matching Alpha Radar member names and emails as an Excel workbook."""
    now = datetime.now(timezone.utc)
    try:
        members = _filter_and_sort_members(
            _load_members(sb, now),
            query=q,
            status_filter=status_filter,
            sort=sort,
        )
        workbook = _build_members_workbook(members)
    except Exception as exc:
        logger.exception("[admin] failed to export Alpha Radar subscribers")
        raise HTTPException(
            status_code=500,
            detail="Unable to create the subscriber export. Please retry shortly.",
        ) from exc

    filename = f"alpha-radar-members-{now.date().isoformat()}.xlsx"
    return Response(
        content=workbook,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
