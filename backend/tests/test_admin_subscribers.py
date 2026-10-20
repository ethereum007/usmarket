from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
import unittest

from openpyxl import load_workbook

from app.routes.admin_subscribers import (
    _build_member,
    _build_members_workbook,
    _matches_search,
    _matches_status,
)


class AdminSubscribersTest(unittest.TestCase):
    def setUp(self) -> None:
        self.now = datetime(2026, 7, 23, tzinfo=timezone.utc)
        self.user = {
            "id": "member-1",
            "email": "member@example.com",
            "phone": "+919876543210",
            "created_at": datetime(2026, 6, 1, tzinfo=timezone.utc),
            "last_sign_in_at": datetime(2026, 7, 22, tzinfo=timezone.utc),
            "user_metadata": {"full_name": "Alpha Member"},
        }

    def test_paid_member_exposes_only_dashboard_fields(self) -> None:
        member = _build_member(
            self.user,
            {
                "plan": "premium",
                "status": "active",
                "razorpay_subscription_id": "order_1234567890",
                "razorpay_customer_id": "pay_1234567890",
                "current_period_start": "2026-07-01T00:00:00+00:00",
                "current_period_end": "2026-07-31T00:00:00+00:00",
            },
            self.now,
        )

        self.assertEqual(member["membership"]["status"], "premium")
        self.assertEqual(member["payment"]["status"], "paid")
        self.assertEqual(member["payment"]["amountPaise"], 99_900)
        self.assertEqual(member["payment"]["paymentReference"], "pay_1…67890")
        self.assertNotIn("user_metadata", member)

    def test_expired_subscription_is_not_reported_as_active_premium(self) -> None:
        member = _build_member(
            self.user,
            {
                "plan": "premium",
                "status": "active",
                "current_period_end": "2026-07-01T00:00:00+00:00",
            },
            self.now,
        )

        self.assertEqual(member["membership"]["status"], "expired")
        self.assertTrue(member["membership"]["isExpired"])
        self.assertTrue(_matches_status(member, "expired"))

    def test_search_uses_only_name_email_and_mobile(self) -> None:
        member = _build_member(self.user, {}, self.now)
        self.assertTrue(_matches_search(member, "member@example.com"))
        self.assertTrue(_matches_search(member, "98765"))
        self.assertFalse(_matches_search(member, "payment-secret"))

    def test_members_workbook_contains_only_name_and_email(self) -> None:
        payload = _build_members_workbook(
            [
                {"name": "Kalyan Kumar", "email": "kalyan@example.com"},
                {"name": "Praneeth Gopu", "email": "praneeth@example.com"},
            ]
        )

        workbook = load_workbook(BytesIO(payload), read_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))
        workbook.close()

        self.assertEqual(rows[0], ("Name", "Email"))
        self.assertEqual(rows[1], ("Kalyan Kumar", "kalyan@example.com"))
        self.assertEqual(rows[2], ("Praneeth Gopu", "praneeth@example.com"))


if __name__ == "__main__":
    unittest.main()
